from pathlib import Path
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'gui'))
from manual_content import MANUAL_SCREENSHOTS, MANUAL_SECTIONS, MANUAL_SECTIONS_EN

# ★操作説明書は日英併記。アプリ内Helpは表示言語で切り替える(screens/manual.py)。
# 日本語章と英語章は同じ順序・同じ項目数で対応している前提。
assert len(MANUAL_SECTIONS) == len(MANUAL_SECTIONS_EN)
for (_ja_title, _ja_items), (_en_title, _en_items) in zip(MANUAL_SECTIONS, MANUAL_SECTIONS_EN):
    assert len(_ja_items) == len(_en_items), (_ja_title, _en_title)
BILINGUAL_SECTIONS = list(zip(MANUAL_SECTIONS, MANUAL_SECTIONS_EN))
EN_COLOR = RGBColor(64, 64, 64)


def add_en_paragraph(text):
    q = d.add_paragraph()
    run = q.add_run(text)
    run.font.color.rgb = EN_COLOR
    return q

OUT=Path(__file__).with_name('shonan_pi5_operation_manual.docx')
IMAGE_DIR=ROOT/'docs'/'images'
d=Document()
sec=d.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
sec.header_distance=sec.footer_distance=Inches(0.492)
sec.different_first_page_header_footer = True

styles=d.styles
normal=styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1.25
for name,size,color,before,after in (
    ('Heading 1',16,'2E74B5',18,10),('Heading 2',13,'2E74B5',14,7),('Heading 3',12,'1F4D78',10,5)):
    s=styles[name];s.font.name='Calibri';s.font.size=Pt(size);s.font.bold=True
    s.font.color.rgb=RGBColor.from_string(color);s.paragraph_format.space_before=Pt(before);s.paragraph_format.space_after=Pt(after)

# editorial_cover
p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(120)
r=p.add_run('Shonan Lite for RasPI5');r.bold=True;r.font.size=Pt(26);r.font.color.rgb=RGBColor(31,77,120)
p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('DVB-S2 DATV送受信システム 操作説明書');r.bold=True;r.font.size=Pt(18)
p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('DVB-S2 DATV Transmit/Receive System Operation Manual');r.bold=True;r.font.size=Pt(14)
p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(24)
p.add_run('対象 / Target: Pi 5 + ADALM-Pluto\n'
          '接続先: Pluto / Pi 5（IPアドレスは設定画面で指定）\n'
          'Connection: Pluto / Pi 5 (IP addresses are set on the Settings screen)\n'
          '版 / Edition: 2026-09-30').font.size=Pt(11)
d.add_page_break()

d.add_heading('目次 / Contents',level=1)
for (title,items),(title_en,items_en) in BILINGUAL_SECTIONS:
    d.add_paragraph(f'{title} / {title_en}',style='List Number')
    for (subtitle,_),(subtitle_en,_) in zip(items,items_en):
        q=d.add_paragraph(f'{subtitle} / {subtitle_en}',style='List Bullet 2');q.paragraph_format.space_after=Pt(2)
d.add_page_break()

# ★クイックスタートの本文はmanual_content.pyの「クイックスタート」章(Helpと共通)を使う。
# 以前はここにも独自の手順(周波数438.200 MHz等の古い内容)を書いており、章が重複していた。
note=d.add_paragraph('本書に記載の内容は、実際に動作させて動作確認した内容です。ただし、まだバグが内在している可能性はあります。\n'
                   'The contents of this manual have been verified by actual operation. However, bugs may still remain.\n'
                   '本書は日本語と英語を併記しています。各項目で日本語の後に英語が続きます。\n'
                   'This manual is written in Japanese and English. In each item, the Japanese text is followed by the English text.')
note.paragraph_format.space_after=Pt(10)
for run in note.runs:
    run.italic=True
    run.font.color.rgb=RGBColor(100,100,100)

for (title,items),(title_en,items_en) in BILINGUAL_SECTIONS:
    d.add_heading(f'{title} / {title_en}',level=1)
    image_names=MANUAL_SCREENSHOTS.get(title, [])
    if isinstance(image_names, str):
        image_names=[image_names]
    for image_name in image_names:
        image_path=IMAGE_DIR/image_name
        if not image_path.exists():
            continue
        p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(image_path), width=Inches(6.0))
        cap=d.add_paragraph(f'図: {title}（実機スクリーンショット） / Figure: {title_en} (device screenshot)')
        cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
        cap.runs[0].italic=True;cap.runs[0].font.size=Pt(9);cap.runs[0].font.color.rgb=RGBColor(100,100,100)
    for (subtitle,text),(subtitle_en,text_en) in zip(items,items_en):
        d.add_heading(f'{subtitle} / {subtitle_en}',level=2)
        d.add_paragraph(text)
        add_en_paragraph(text_en)

d.add_heading('付録A 表示メッセージ早見表 / Appendix A Quick Reference of Displayed Messages',level=1)
t=d.add_table(rows=1,cols=2);t.style='Table Grid';t.autofit=False
t.columns[0].width=Inches(1.875);t.columns[1].width=Inches(4.625)
for c,v in zip(t.rows[0].cells,('表示 / Display','意味・対処 / Meaning and action')):c.text=v
for a,b,b_en in (
    ('IIO preflight OK','RX DMAから4,096 byte取得成功。受信処理を開始できます。',
     'Read 4,096 bytes from RX DMA successfully. Reception can start.'),
    ('IIO preflight FAILED','RF NO LOCKではなくIIO/RX DMA異常。Pluto再起動と接続確認を行います。',
     'Not RF NO LOCK but an IIO/RX DMA fault. Reboot the Pluto and check the connection.'),
    ('Sync: Locked','DVB-S2同期成功。TS packets増加も確認します。',
     'DVB-S2 sync succeeded. Also check that TS packets are increasing.'),
    ('Lost / Rate','TS continuity推定ロス数 / 直近の受信ビットレート。Lostの急増はTS欠落を示します。',
     'Estimated losses from TS continuity / recent receive bit rate. A sudden rise in Lost indicates missing TS.'),
    ('Video Waiting','SPS/PPSと最初のIDRを待っています。',
     'Waiting for SPS/PPS and the first IDR.'),
    ('Video Ready (IDR)','SPS/PPS後のIDRを検出し、映像表示可能です。',
     'An IDR after SPS/PPS was detected; video can be displayed.'),
):
    cells=t.add_row().cells;cells[0].text=a;cells[1].text=f'{b}\n{b_en}'

first_footer=sec.first_page_footer.paragraphs[0]
first_footer.alignment=WD_ALIGN_PARAGRAPH.LEFT
first_footer.paragraph_format.line_spacing = 1.0
credit=first_footer.add_run(
    '受信部の方式考案・受信部原システム設計: 山崎慎慈氏(JE1BTA) '
    'rpi-dvbs2-receiver-guiの設計に基づきます\n'
    '受信部安定化調査修正・再捕捉修正・本アプリ開発: 真城和一\n'
    'Reception method and original reception system design: Shinji Yamazaki (JE1BTA), '
    'based on the design of rpi-dvbs2-receiver-gui\n'
    'Reception stability investigation and fixes, re-acquisition fixes, and development of this app: Kazuichi Shinjo'
)
credit.font.size=Pt(8)
credit.font.color.rgb=RGBColor(100,100,100)

header=sec.header.paragraphs[0];header.text='Shonan Lite for RasPI5 操作説明書 / Operation Manual';header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('Page ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

d.save(OUT)
print(OUT)
